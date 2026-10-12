"""Ayudas compartidas de los tests del censo: planillas de prueba y sectores reales.

Las planillas se arman con openpyxl, como la real: encabezados en la fila 1, código como
fórmula COUNTIF sin valor guardado y listas de catálogo en las validaciones de datos.
Nunca se usa el archivo real.
"""

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from django.contrib.auth.models import Group, User
from django.test import Client
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.censo.models import Arbol, Brigada, Especie, GrupoBrigada, Sector, ValorCatalogo
from apps.censo.servicios.calidad_registro import aplicar
from apps.censo.servicios.importar_sectores import importar_sectores
from apps.censo.servicios.resolver_campos import Resolutor, aplicar_textos
from config.constantes.catalogos import CAMPO_ORIGEN, CAMPOS_ESTADO_ARBOL, CAMPOS_TEXTO_RESUELTOS
from config.constantes.permisos import NOMBRE_ROL_ADMINISTRADOR, NOMBRE_ROL_EDITOR
from config.constantes.planilla import (
    COLUMNAS_OPCIONALES_PLANILLA,
    COLUMNAS_PLANILLA,
    HOJA_PLANILLA,
)

RUTA_GEOJSON = (
    Path(__file__).resolve().parents[2] / "datos" / "sectores" / "sectores-utm18s-v0.geojson"
)
CATALOGOS = {
    "estado_general": "Bueno,Regular,Malo",
    "estado_copa": "Completa,Parcial,Escasa",
    "origen": "Nativa,Introducida",
}
FILAS_CON_VALIDACION = 3000
NOMBRE_ARCHIVO = "arb-prueba.xlsx"
FORMULA_CODIGO = '=IF(B{n}="","","S"&TEXT(B{n},"00")&"-A"&TEXT(COUNTIF(B$2:B{n},B{n}),"000"))'

# Valores de una fila completa (nombres de campo del modelo), usados por el test de fila completa.
TEXTOS_FILA_COMPLETA = {
    "tipo_individuo": "Árbol",
    "nombre_comun": "Molle",
    "nombre_cientifico": "Schinus molle",
    "identificacion": "Confirmada",
    "origen": "Nativa",
    "condicion_conservacion": "Sin categoría",
    "estado_general": "Bueno",
    "estado_copa": "Completa",
    "tronco_danos": "Sin daños",
    "raices_base": "Sanas",
    "interferencia_entorno": "Ninguna",
    "foto_archivo": "IMG_0001.jpg",
    "foto_detalle_archivo": "IMG_0002.jpg",
    "observaciones": "Junto a la banca",
}
MEDIDAS_FILA_COMPLETA = {
    "altura_total_m": 8.5,
    "dap_cm": 35.25,
    "copa_ns_m": 4.5,
    "copa_eo_m": 5,
    "co2_almacenado_kg": 120.5,
    "co2_captura_anual_kg": 10.25,
}


def crear_planilla(ruta, filas, sin_columnas=(), columnas_ajenas=(), renombrar=None):
    """Escribe una planilla de prueba; cada fila es {campo: valor}.

    Sin `codigo` en la fila se escribe la fórmula de la planilla real (sin valor guardado). Trae
    las columnas opcionales (salvo las de `sin_columnas`) y, al final, las `columnas_ajenas`:
    títulos que la planilla no usa y que se deben ignorar. Con `renombrar` (`{encabezado: otro}`)
    una columna se escribe con otro título, por ejemplo Latitud donde va «UTM Este (m)».
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA_PLANILLA
    columnas = {**COLUMNAS_PLANILLA, **COLUMNAS_OPCIONALES_PLANILLA}
    encabezados = [enc for enc in columnas if enc not in sin_columnas]
    renombrar = renombrar or {}
    hoja.append([*(renombrar.get(enc, enc) for enc in encabezados), *columnas_ajenas])
    columna_por_campo = {columnas[enc]: indice for indice, enc in enumerate(encabezados, start=1)}
    for numero_fila, fila in enumerate(filas, start=2):
        for campo, valor in fila.items():
            if campo in columna_por_campo:
                hoja.cell(row=numero_fila, column=columna_por_campo[campo], value=valor)
        if "codigo" not in fila and "codigo" in columna_por_campo:
            hoja.cell(
                row=numero_fila,
                column=columna_por_campo["codigo"],
                value=FORMULA_CODIGO.format(n=numero_fila),
            )
    for campo, lista in CATALOGOS.items():
        if campo in columna_por_campo:
            letra = get_column_letter(columna_por_campo[campo])
            validacion = DataValidation(type="list", formula1=f'"{lista}"')
            validacion.add(f"{letra}2:{letra}{FILAS_CON_VALIDACION + 1}")
            hoja.add_data_validation(validacion)
    libro.save(ruta)
    return ruta


@pytest.fixture
def sectores(db):
    """Los 5 sectores reales del GeoJSON del repo."""
    importar_sectores(RUTA_GEOJSON)


@pytest.fixture
def punto_dentro(sectores):
    """Coordenadas UTM de un punto dentro del sector 1."""
    punto = Sector.objects.get(id=1).geom.point_on_surface
    return punto.x, punto.y


@pytest.fixture
def ruta_planilla(tmp_path):
    return tmp_path / NOMBRE_ARCHIVO


def fila(**extra):
    datos = {"sector": 1, "brigada": "Brigada A"}
    datos.update(extra)
    return datos


FUENTE_PANEL = "Prueba del panel de administración"
CLAVE_PRUEBA = "clave-de-prueba-1234"


def arbol_con_textos(codigo, sector_id, **datos):
    """Un árbol sin guardar. Los nombres, el origen, los estados y la brigada se dan como texto.

    Se resuelven con el mismo código de producción. La especie se crea si no existe (con el origen
    del catálogo, si ya está cargado); lo que no se resuelve queda en `valores_originales`.
    """
    textos = {campo: datos.pop(campo) for campo in CAMPOS_TEXTO_RESUELTOS if campo in datos}
    arbol = Arbol(codigo=codigo, sector_id=sector_id, **datos)
    _asegurar_especie(textos)
    aplicar_textos(arbol, textos)
    return arbol


def arbol_en_memoria_con_textos(codigo, sector_id, **datos):
    """Un árbol sin guardar ni consultar la BD: cada texto va en una fila suelta, sin guardar.

    Sirve para probar la regla de calidad, que solo lee el texto de los datos.
    """
    textos = {campo: datos.pop(campo) for campo in CAMPOS_TEXTO_RESUELTOS if campo in datos}
    estados = {
        campo: ValorCatalogo(campo=campo, valor=textos[campo])
        for campo in CAMPOS_ESTADO_ARBOL
        if textos.get(campo) is not None
    }
    arbol = Arbol(codigo=codigo, sector_id=sector_id, **datos, **estados)
    nombres = [textos.get(c) for c in ("nombre_cientifico", "nombre_comun", CAMPO_ORIGEN)]
    if any(nombre is not None for nombre in nombres):
        origen = textos.get(CAMPO_ORIGEN)
        arbol.especie = Especie(
            nombre_cientifico=textos.get("nombre_cientifico"),
            nombre_comun=textos.get("nombre_comun"),
            origen=None if origen is None else ValorCatalogo(campo=CAMPO_ORIGEN, valor=origen),
        )
    if textos.get("brigada") is not None:
        arbol.grupo = GrupoBrigada(nombre="grupo", brigada=Brigada(nombre=textos["brigada"]))
    return arbol


def _asegurar_especie(textos):
    """Crea la especie del nombre científico si no hay una que coincida."""
    cientifico = textos.get("nombre_cientifico")
    if not cientifico or Resolutor.desde_bd().especie(cientifico) is not None:
        return
    origen = ValorCatalogo.objects.filter(campo=CAMPO_ORIGEN, valor=textos.get(CAMPO_ORIGEN))
    Especie.objects.create(
        nombre_cientifico=cientifico, nombre_comun=textos.get("nombre_comun"), origen=origen.first()
    )


@pytest.fixture
def especie_molle(db):
    """La especie de las planillas de prueba: la ingesta nunca crea una especie que no existe."""
    return Especie.objects.create(nombre_cientifico="Schinus molle", nombre_comun="Molle")


@pytest.fixture
def especies_demo(especie_molle):
    """Las tres especies de los árboles simulados (con Schinus molle de `especie_molle`)."""
    Especie.objects.create(nombre_cientifico="Prosopis pallida", nombre_comun="Huarango")
    Especie.objects.create(nombre_cientifico="Roystonea regia", nombre_comun="Palma real")


@pytest.fixture
def nuevo_arbol(sectores):
    """Fábrica de árboles de prueba: siempre con fuente y `demo=True` (nunca datos del censo)."""

    def crear(codigo="S01-A001", sector_id=1, **extra):
        datos = {"fuente": FUENTE_PANEL, "demo": True}
        datos.update(extra)
        arbol = arbol_con_textos(codigo, sector_id, **datos)
        arbol.save()
        return arbol

    return crear


VALORES_COMPLETOS = {
    "fecha_registro": date(2026, 10, 1),
    "brigada": "Brigada A",
    "tipo_individuo": "Árbol ≥1.30 m (medir DAP)",
    "nombre_comun": "Molle",
    "nombre_cientifico": "Schinus molle",
    "identificacion": "Identificada",
    "origen": "Nativa",
    "condicion_conservacion": "No evaluada",
    "altura_total_m": Decimal("8.5"),
    "dap_cm": Decimal("35.25"),
    "copa_ns_m": Decimal("4.5"),
    "copa_eo_m": Decimal("5"),
    "co2_almacenado_kg": Decimal("120.5"),
    "co2_captura_anual_kg": Decimal("10.25"),
    "estado_general": "Bueno",
    "estado_copa": "Buena",
    "tronco_danos": "Sin daño aparente",
    "raices_base": "Sin daño aparente",
    "interferencia_entorno": "Ninguna",
}
"""Un árbol completo y consistente (valores del catálogo de las cápsulas), sin ubicación."""


def valores_de_arbol_completo():
    """Copia de los valores de un árbol completo, lista para cambiar."""
    return dict(VALORES_COMPLETOS)


def arbol_evaluado(codigo="S01-A001", sector_id=1, demo=False, **extra):
    """Árbol de prueba completo y pasado por la regla de calidad: sale en la API.

    Por defecto es de campo (`demo=False`, con brigada); para uno de ejemplo, `demo=True` y
    `brigada=None`. Su punto está dentro de su sector. Con `extra` se cambia cualquier campo
    (y la regla vuelve a decidir el estado). Siempre lleva `fuente`.
    """
    datos = valores_de_arbol_completo()
    datos.update(
        ubicacion=Sector.objects.get(id=sector_id).geom.point_on_surface,
        fuente=FUENTE_PANEL,
        demo=demo,
    )
    datos.update(extra)
    arbol = arbol_con_textos(codigo, sector_id, **datos)
    aplicar(arbol)
    arbol.save()
    return arbol


def arbol_de_ejemplo(codigo="S01-A001", sector_id=1, **extra):
    """Como `arbol_evaluado`, pero de ejemplo (`demo=True`, sin brigada)."""
    extra.setdefault("brigada", None)
    return arbol_evaluado(codigo, sector_id, demo=True, **extra)


def usuario_con_rol(nombre_usuario, nombre_rol):
    """Usuario del panel (con acceso al admin) dentro del grupo que creó la migración de roles."""
    usuario = User.objects.create_user(nombre_usuario, password=CLAVE_PRUEBA, is_staff=True)
    usuario.groups.add(Group.objects.get(name=nombre_rol))
    return usuario


@pytest.fixture
def usuario_administrador(db):
    return usuario_con_rol("administrador-prueba", NOMBRE_ROL_ADMINISTRADOR)


@pytest.fixture
def usuario_editor(db):
    return usuario_con_rol("editor-prueba", NOMBRE_ROL_EDITOR)


def cliente_de(usuario):
    """Cliente del admin con la sesión de ese usuario."""
    cliente = Client()
    cliente.force_login(usuario)
    return cliente


@pytest.fixture
def cliente_administrador(usuario_administrador):
    return cliente_de(usuario_administrador)


@pytest.fixture
def cliente_editor(usuario_editor):
    return cliente_de(usuario_editor)
