"""Importa la planilla de la brigada (hoja Registro_Campo). Se puede repetir sin duplicar."""

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.contrib.gis.geos import Point
from django.db import models, transaction
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from apps.censo.models import Arbol, Sector
from config.constantes import (
    CAMPOS_CON_NOTA_PLANILLA,
    CAMPOS_NUMERICOS_PLANILLA,
    COLUMNAS_PLANILLA,
    FILA_ENCABEZADOS_PLANILLA,
    FORMATO_CODIGO_CALCULADO,
    FORMATO_NOTA_PLANILLA,
    FORMATOS_FECHA_PLANILLA,
    HOJA_PLANILLA,
    PATRON_CODIGO_ARBOL,
    SEPARADOR_NOTAS,
    SRID_UTM_CENSO,
)

CAMPOS_SIN_COPIA_DIRECTA = frozenset(
    {"codigo", "sector", "fecha_registro", "observaciones", "control_registro"}
)
"""Campos con tratamiento propio; `control_registro` se lee pero no se guarda."""

CAMPOS_TEXTO = (
    frozenset(COLUMNAS_PLANILLA.values()) - CAMPOS_NUMERICOS_PLANILLA - CAMPOS_SIN_COPIA_DIRECTA
)
"""Campos de texto que se copian tal cual (sin espacios en los bordes) al modelo."""

CAMPOS_DECIMALES = {
    campo.name: campo
    for campo in Arbol._meta.get_fields()
    if isinstance(campo, models.DecimalField)
}
"""Campos decimales del modelo: de ellos salen los límites de dígitos (`max_digits`, etc.)."""

ENCABEZADO_POR_CAMPO = {campo: encabezado for encabezado, campo in COLUMNAS_PLANILLA.items()}
"""Inverso de COLUMNAS_PLANILLA, para nombrar la columna en los reportes."""


@dataclass
class ValorNoValido:
    """Un valor de la planilla que no se pudo leer como se esperaba."""

    codigo: str
    columna: str
    texto: str


@dataclass
class ArbolFueraDeSector:
    """Un árbol cuyo punto cae fuera del polígono de su sector."""

    codigo: str
    sector: int
    distancia_m: float


@dataclass
class ReportePlanilla:
    """Resultado de importar la planilla (el mismo con y sin `--seco`)."""

    archivo: str = ""
    seco: bool = False
    filas_leidas: int = 0
    creados: int = 0
    actualizados: int = 0
    sin_coordenadas: int = 0
    codigos_calculados: int = 0
    numeros_no_validos: list[ValorNoValido] = field(default_factory=list)
    fechas_no_validas: list[ValorNoValido] = field(default_factory=list)
    fuera_de_catalogo: list[ValorNoValido] = field(default_factory=list)
    fuera_de_sector: list[ArbolFueraDeSector] = field(default_factory=list)


@dataclass
class _Contexto:
    """Lo que comparten todas las filas de una importación."""

    nombre_archivo: str
    sectores: dict
    catalogos: dict
    reporte: ReportePlanilla
    orden_en_sector: dict = field(default_factory=dict)
    codigos_vistos: set = field(default_factory=set)


def importar_planilla(ruta, seco=False):
    """Lee la planilla y hace upsert por `codigo`. Con `seco` valida y revierte al final.

    Lanza ValueError con un mensaje claro si el archivo no sirve.
    """
    ruta = Path(ruta)
    if not ruta.is_file():
        raise ValueError(f"No existe el archivo: {ruta}")
    hoja_formulas, hoja_valores = _abrir_hojas(ruta)
    columnas = _ubicar_columnas(hoja_formulas)
    filas = _leer_filas(hoja_formulas, hoja_valores, columnas)
    reporte = ReportePlanilla(archivo=ruta.name, seco=seco, filas_leidas=len(filas))
    contexto = _Contexto(
        nombre_archivo=ruta.name,
        sectores=_cargar_sectores(filas),
        catalogos=_leer_catalogos(hoja_formulas, columnas),
        reporte=reporte,
    )
    with transaction.atomic():
        for numero_fila, fila in filas:
            _importar_fila(numero_fila, fila, contexto)
        if seco:
            transaction.set_rollback(True)
    return reporte


def _abrir_hojas(ruta):
    """Abre el libro dos veces: con fórmulas (y validaciones) y con los valores guardados."""
    try:
        libro_formulas = load_workbook(ruta, data_only=False)
        libro_valores = load_workbook(ruta, data_only=True)
    except Exception as error:  # openpyxl lanza tipos distintos según qué esté roto
        raise ValueError(f"No se pudo abrir {ruta.name} como planilla de Excel: {error}") from error
    if HOJA_PLANILLA not in libro_formulas.sheetnames:
        raise ValueError(f"La planilla no tiene la hoja «{HOJA_PLANILLA}».")
    return libro_formulas[HOJA_PLANILLA], libro_valores[HOJA_PLANILLA]


def _ubicar_columnas(hoja):
    """Devuelve {campo: número de columna} buscando los encabezados por nombre."""
    posicion_por_encabezado = {
        str(celda.value).strip(): celda.column
        for celda in hoja[FILA_ENCABEZADOS_PLANILLA]
        if celda.value is not None
    }
    faltan = [
        encabezado for encabezado in COLUMNAS_PLANILLA if encabezado not in posicion_por_encabezado
    ]
    if faltan:
        raise ValueError("Faltan columnas en la planilla: " + ", ".join(faltan))
    return {campo: posicion_por_encabezado[enc] for enc, campo in COLUMNAS_PLANILLA.items()}


def _leer_catalogos(hoja, columnas):
    """Lee las listas de las validaciones de datos: {campo: {valores permitidos}}."""
    campo_por_columna = {columna: campo for campo, columna in columnas.items()}
    catalogos = {}
    for validacion in hoja.data_validations.dataValidation:
        formula = (validacion.formula1 or "").strip()
        if validacion.type != "list" or not (formula.startswith('"') and formula.endswith('"')):
            continue  # solo se leen las listas escritas dentro de la validación
        valores = {valor.strip() for valor in formula.strip('"').split(",") if valor.strip()}
        for rango in str(validacion.sqref).split():
            primera_columna, _, ultima_columna, _ = range_boundaries(rango)
            for columna in range(primera_columna, ultima_columna + 1):
                if columna in campo_por_columna:
                    catalogos.setdefault(campo_por_columna[columna], set()).update(valores)
    return catalogos


def _leer_filas(hoja_formulas, hoja_valores, columnas):
    """Devuelve [(número de fila, {campo: valor guardado})] de las filas que tienen Sector."""
    filas = []
    for numero_fila in range(FILA_ENCABEZADOS_PLANILLA + 1, hoja_formulas.max_row + 1):
        fila = {
            campo: hoja_valores.cell(row=numero_fila, column=columna).value
            for campo, columna in columnas.items()
        }
        if _texto(fila["sector"]) is not None:
            filas.append((numero_fila, fila))
    return filas


def _cargar_sectores(filas):
    """Trae de la BD los sectores que usa la planilla; falla si falta alguno."""
    ids = {_entero_sector(fila["sector"]) for _, fila in filas}
    sectores = Sector.objects.in_bulk(ids)
    faltan = sorted(ids - set(sectores))
    if faltan:
        raise ValueError(
            "La planilla usa sectores que no están en la BD: "
            + ", ".join(str(id_sector) for id_sector in faltan)
            + ". Importa antes los sectores (importar_sectores)."
        )
    return sectores


def _importar_fila(numero_fila, fila, contexto):
    """Valida una fila y hace upsert del árbol; todo hallazgo va al reporte."""
    reporte = contexto.reporte
    id_sector = _entero_sector(fila["sector"])
    orden = contexto.orden_en_sector.get(id_sector, 0) + 1
    contexto.orden_en_sector[id_sector] = orden
    codigo = _texto(fila["codigo"])
    if codigo is None:
        codigo = FORMATO_CODIGO_CALCULADO.format(sector=id_sector, orden=orden)
        reporte.codigos_calculados += 1
    if not re.fullmatch(PATRON_CODIGO_ARBOL, codigo):
        raise ValueError(f"Fila {numero_fila}: el código «{codigo}» no cumple el patrón.")
    if codigo in contexto.codigos_vistos:
        raise ValueError(f"Fila {numero_fila}: el código «{codigo}» está repetido en la planilla.")
    contexto.codigos_vistos.add(codigo)

    valores = {campo: _texto(fila[campo]) for campo in CAMPOS_TEXTO}
    for campo in valores.keys() & contexto.catalogos.keys():
        if valores[campo] is not None and valores[campo] not in contexto.catalogos[campo]:
            reporte.fuera_de_catalogo.append(
                ValorNoValido(codigo, ENCABEZADO_POR_CAMPO[campo], valores[campo])
            )
    notas = []
    for campo in (c for c in COLUMNAS_PLANILLA.values() if c in CAMPOS_NUMERICOS_PLANILLA):
        valores[campo], texto = _numero(fila[campo], campo)
        if texto is not None:
            encabezado = ENCABEZADO_POR_CAMPO[campo]
            reporte.numeros_no_validos.append(ValorNoValido(codigo, encabezado, texto))
            if campo in CAMPOS_CON_NOTA_PLANILLA:
                notas.append(FORMATO_NOTA_PLANILLA.format(encabezado=encabezado, texto=texto))
    valores["fecha_registro"], texto_fecha = _fecha(fila["fecha_registro"])
    if texto_fecha is not None:
        reporte.fechas_no_validas.append(
            ValorNoValido(codigo, ENCABEZADO_POR_CAMPO["fecha_registro"], texto_fecha)
        )
    partes_observaciones = [_texto(fila["observaciones"]), *notas]
    valores["observaciones"] = SEPARADOR_NOTAS.join(p for p in partes_observaciones if p) or None

    este, norte = valores.pop("utm_este_m"), valores.pop("utm_norte_m")
    valores["ubicacion"] = None
    if este is not None and norte is not None:
        valores["ubicacion"] = Point(float(este), float(norte), srid=SRID_UTM_CENSO)
    else:
        reporte.sin_coordenadas += 1
    sector = contexto.sectores[id_sector]
    if valores["ubicacion"] is not None and not sector.geom.covers(valores["ubicacion"]):
        distancia_m = sector.geom.distance(valores["ubicacion"])
        reporte.fuera_de_sector.append(ArbolFueraDeSector(codigo, id_sector, distancia_m))

    valores["sector"] = sector
    origen = f"hoja {HOJA_PLANILLA} de {contexto.nombre_archivo}"
    partes_fuente = [valores["brigada"], origen, f"fila {numero_fila}"]
    valores["fuente"] = ", ".join(parte for parte in partes_fuente if parte)
    valores["fila_origen"] = numero_fila
    valores["demo"] = False
    _, creado = Arbol.objects.update_or_create(codigo=codigo, defaults=valores)
    if creado:
        reporte.creados += 1
    else:
        reporte.actualizados += 1


def _entero_sector(valor):
    try:
        return int(Decimal(str(valor).strip()))
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"Sector no válido en la planilla: «{valor}».") from error


def _texto(valor):
    """Texto sin espacios en los bordes; None si está vacío."""
    if valor is None:
        return None
    return str(valor).strip() or None


def _numero(valor, campo):
    """Devuelve (Decimal o None, texto no válido o None). Vacío es None sin aviso.

    No es válido lo que no es número ni cabe en el campo decimal del modelo.
    """
    texto = _texto(valor)
    if texto is None:
        return None, None
    try:
        numero = Decimal(texto.replace(",", "."))
    except InvalidOperation:
        return None, texto
    if isinstance(valor, bool) or not numero.is_finite() or not _cabe(numero, campo):
        return None, texto
    return numero, None


def _cabe(numero, campo):
    """True si el número cabe en el DecimalField del modelo (los campos sin modelo, siempre)."""
    campo_modelo = CAMPOS_DECIMALES.get(campo)
    if campo_modelo is None:
        return True
    try:
        redondeado = numero.quantize(Decimal(1).scaleb(-campo_modelo.decimal_places))
    except InvalidOperation:
        return False
    return len(redondeado.as_tuple().digits) <= campo_modelo.max_digits


def _fecha(valor):
    """Devuelve (date o None, texto no válido o None). Vacío es None sin aviso."""
    if isinstance(valor, datetime):
        return valor.date(), None
    if isinstance(valor, date):
        return valor, None
    texto = _texto(valor)
    if texto is None:
        return None, None
    for formato in FORMATOS_FECHA_PLANILLA:
        try:
            return datetime.strptime(texto, formato).date(), None
        except ValueError:
            continue
    return None, texto
