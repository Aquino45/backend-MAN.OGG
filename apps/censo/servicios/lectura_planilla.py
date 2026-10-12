"""Lee la hoja de la brigada (.xlsx) y convierte sus valores; no toca la BD ni evalúa nada.

Solo mira los valores guardados: no se evalúa ninguna fórmula ni macro.
"""

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from django.contrib.gis.geos import Point
from django.db import models
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from apps.censo.models import Arbol
from config.constantes import (
    AVISO_COLUMNAS_COORDENADAS_CON_UTM,
    CAMPO_COORDENADAS_NO_RECONOCIDAS,
    COLUMNAS_OPCIONALES_PLANILLA,
    COLUMNAS_PLANILLA,
    ENCABEZADOS_LATITUD_PLANILLA,
    ENCABEZADOS_LONGITUD_PLANILLA,
    FILA_ENCABEZADOS_PLANILLA,
    FORMATO_COORDENADAS_NO_RECONOCIDAS,
    FORMATOS_FECHA_PLANILLA,
    HOJA_PLANILLA,
    MENSAJE_CARGA_MUCHAS_FILAS,
    RANGO_LATITUD_CAMPUS_GRADOS,
    RANGO_LONGITUD_CAMPUS_GRADOS,
    RANGO_UTM_ESTE_M,
    RANGO_UTM_NORTE_M,
    SRID_MAPA,
    SRID_UTM_CENSO,
    TEXTO_CELDA_VACIA_COORDENADAS,
)

CAMPOS_DECIMALES = {
    campo.name: campo
    for campo in Arbol._meta.get_fields()
    if isinstance(campo, models.DecimalField)
}
"""Campos decimales del modelo: de ellos salen los límites de dígitos (`max_digits`, etc.)."""

ENCABEZADO_POR_CAMPO = {
    campo: encabezado
    for encabezado, campo in {**COLUMNAS_PLANILLA, **COLUMNAS_OPCIONALES_PLANILLA}.items()
}
"""Inverso de COLUMNAS_PLANILLA y las opcionales, para nombrar la columna en los reportes."""


CAMPOS_DE_COORDENADAS = frozenset({"utm_este_m", "utm_norte_m"})
ENCABEZADOS_UTM = tuple(
    encabezado for encabezado, campo in COLUMNAS_PLANILLA.items() if campo in CAMPOS_DE_COORDENADAS
)
"""Los campos y los encabezados de las coordenadas UTM de la planilla."""


@dataclass
class PlanillaLeida:
    """Las filas de la hoja como vienen: `[(número de fila, {campo: valor})]`.

    `avisos` son los avisos de toda la planilla (no de una fila), por ejemplo que las columnas
    Latitud y Longitud traían valores UTM.
    """

    nombre_archivo: str
    filas: list
    catalogos: dict = field(default_factory=dict)
    avisos: list = field(default_factory=list)


def leer(origen, nombre_archivo=None, filas_maximas=None):
    """Lee la hoja de la brigada desde una ruta o desde bytes (`bytes` o `BytesIO`).

    Solo mira los valores guardados: no se evalúa ninguna fórmula ni macro. Lanza ValueError con un
    mensaje claro si el archivo no sirve o declara más de `filas_maximas` filas.
    """
    if isinstance(origen, (str, Path)):
        ruta = Path(origen)
        if not ruta.is_file():
            raise ValueError(f"No existe el archivo: {ruta}")
        contenido, nombre = ruta.read_bytes(), nombre_archivo or ruta.name
    else:
        contenido = origen.getvalue() if isinstance(origen, BytesIO) else bytes(origen)
        nombre = nombre_archivo or ""
    hoja_formulas, hoja_valores = _abrir_hojas(contenido, nombre)
    declaradas = hoja_formulas.max_row - FILA_ENCABEZADOS_PLANILLA
    if filas_maximas is not None and declaradas > filas_maximas:
        raise ValueError(MENSAJE_CARGA_MUCHAS_FILAS.format(filas=declaradas, maximo=filas_maximas))
    columnas = _ubicar_columnas(hoja_formulas)
    filas = _leer_filas(hoja_formulas, hoja_valores, columnas)
    return PlanillaLeida(
        nombre_archivo=nombre,
        filas=filas,
        catalogos=_leer_catalogos(hoja_formulas, columnas),
        avisos=_interpretar_coordenadas(filas, columnas),
    )


def _abrir_hojas(contenido, nombre):
    """Abre el libro dos veces: con fórmulas (y validaciones) y con los valores guardados."""
    try:
        libro_formulas = load_workbook(BytesIO(contenido), data_only=False)
        libro_valores = load_workbook(BytesIO(contenido), data_only=True)
    except Exception as error:  # openpyxl lanza tipos distintos según qué esté roto
        raise ValueError(f"No se pudo abrir {nombre} como planilla de Excel: {error}") from error
    if HOJA_PLANILLA not in libro_formulas.sheetnames:
        raise ValueError(f"La planilla no tiene la hoja «{HOJA_PLANILLA}».")
    return libro_formulas[HOJA_PLANILLA], libro_valores[HOJA_PLANILLA]


def _ubicar_columnas(hoja):
    """Devuelve {campo: número de columna} buscando los encabezados por nombre.

    Las columnas opcionales entran solo si la planilla las trae; las ajenas se ignoran. En lugar de
    «UTM Este (m)» y «UTM Norte (m)» la planilla puede traer el par Latitud y Longitud: entran como
    los campos `latitud` y `longitud`, que `_interpretar_coordenadas` convierte.
    """
    posicion_por_encabezado = {
        str(celda.value).strip(): celda.column
        for celda in hoja[FILA_ENCABEZADOS_PLANILLA]
        if celda.value is not None
    }
    par_alternativo = _par_de_coordenadas(posicion_por_encabezado)
    faltan = [
        encabezado
        for encabezado, campo in COLUMNAS_PLANILLA.items()
        if encabezado not in posicion_por_encabezado
        and not (par_alternativo and campo in CAMPOS_DE_COORDENADAS)
    ]
    if faltan:
        raise ValueError("Faltan columnas en la planilla: " + ", ".join(faltan))
    columnas = {
        campo: posicion_por_encabezado[encabezado]
        for encabezado, campo in {**COLUMNAS_PLANILLA, **COLUMNAS_OPCIONALES_PLANILLA}.items()
        if encabezado in posicion_por_encabezado
    }
    return {**columnas, **(par_alternativo or {})}


def _par_de_coordenadas(posicion_por_encabezado):
    """`{latitud: columna, longitud: columna}` si la planilla trae el par sin UTM; si no, None."""
    if any(encabezado in posicion_por_encabezado for encabezado in ENCABEZADOS_UTM):
        return None
    par = {}
    for campo, encabezados in (
        ("latitud", ENCABEZADOS_LATITUD_PLANILLA),
        ("longitud", ENCABEZADOS_LONGITUD_PLANILLA),
    ):
        presentes = [e for e in encabezados if e in posicion_por_encabezado]
        if presentes:
            par[campo] = posicion_por_encabezado[presentes[0]]
    return par if len(par) == 2 else None


def _leer_catalogos(hoja, columnas):
    """Lee las listas de las validaciones de datos: {campo: {valores permitidos}}."""
    campo_por_columna = {columna: campo for campo, columna in columnas.items()}
    catalogos = {}
    for validacion in hoja.data_validations.dataValidation:
        formula = (validacion.formula1 or "").strip()
        if validacion.type != "list" or not (formula.startswith('"') and formula.endswith('"')):
            continue  # solo se leen las listas escritas dentro de la validación
        valores = {valor.strip() for valor in formula.strip('"').split(",") if valor.strip()}
        for rango in str(validacion.sqref).split():
            primera_columna, _, ultima_columna, _ = range_boundaries(rango)
            for columna in range(primera_columna, ultima_columna + 1):
                if columna in campo_por_columna:
                    catalogos.setdefault(campo_por_columna[columna], set()).update(valores)
    return catalogos


def _leer_filas(hoja_formulas, hoja_valores, columnas):
    """Devuelve [(número de fila, {campo: valor guardado})] de las filas que tienen Sector."""
    filas = []
    for numero_fila in range(FILA_ENCABEZADOS_PLANILLA + 1, hoja_formulas.max_row + 1):
        fila = {
            campo: hoja_valores.cell(row=numero_fila, column=columna).value
            for campo, columna in columnas.items()
        }
        if texto_de(fila["sector"]) is not None:
            filas.append((numero_fila, fila))
    return filas


# --- coordenadas en grados o con el rótulo equivocado ------------------------------------------


def _interpretar_coordenadas(filas, columnas):
    """Pasa las columnas Latitud y Longitud a `utm_este_m` y `utm_norte_m`; devuelve los avisos.

    Cada fila se lee sola: los grados WGS 84 del campus se convierten a UTM 18S; los valores de 6
    y 7 cifras se toman como UTM (el rótulo está mal) y la planilla avisa una vez; otra cosa deja
    la fila sin coordenadas, con su texto en `coordenadas_no_reconocidas`, para que vaya a revisión.
    """
    if "latitud" not in columnas:
        return []
    hubo_utm_con_rotulo_de_grados = False
    for _, datos in filas:
        primera, segunda = datos.pop("latitud"), datos.pop("longitud")
        este, norte, eran_utm = _coordenadas_utm_de(primera, segunda)
        datos["utm_este_m"], datos["utm_norte_m"] = este, norte
        hubo_utm_con_rotulo_de_grados |= eran_utm
        if este is None and (texto_de(primera) or texto_de(segunda)):
            datos[CAMPO_COORDENADAS_NO_RECONOCIDAS] = FORMATO_COORDENADAS_NO_RECONOCIDAS.format(
                primera=texto_de(primera) or TEXTO_CELDA_VACIA_COORDENADAS,
                segunda=texto_de(segunda) or TEXTO_CELDA_VACIA_COORDENADAS,
            )
    return [AVISO_COLUMNAS_COORDENADAS_CON_UTM] if hubo_utm_con_rotulo_de_grados else []


def _coordenadas_utm_de(primera, segunda):
    """`(este, norte, eran_utm)` de lo que traen las columnas Latitud y Longitud.

    Si no se reconocen (o falta alguna) devuelve `(None, None, False)`.
    """
    en_primera, en_segunda = _decimal_de(primera), _decimal_de(segunda)
    if en_primera is None or en_segunda is None:
        return None, None, False
    if _en_rango(en_primera, RANGO_LATITUD_CAMPUS_GRADOS) and _en_rango(
        en_segunda, RANGO_LONGITUD_CAMPUS_GRADOS
    ):
        punto = Point(float(en_segunda), float(en_primera), srid=SRID_MAPA)
        punto.transform(SRID_UTM_CENSO)
        return punto.x, punto.y, False
    if _en_rango(en_primera, RANGO_UTM_ESTE_M, True) and _en_rango(
        en_segunda, RANGO_UTM_NORTE_M, True
    ):
        return primera, segunda, True
    return None, None, False


def _decimal_de(valor):
    """El valor como Decimal (con coma o punto decimal), o None si está vacío o no es un número."""
    texto = texto_de(valor)
    if texto is None or isinstance(valor, bool):
        return None
    try:
        numero = Decimal(texto.replace(",", "."))
    except InvalidOperation:
        return None
    return numero if numero.is_finite() else None


def _en_rango(numero, rango, excluye_maximo=False):
    minimo, maximo = rango
    return minimo <= numero < maximo if excluye_maximo else minimo <= numero <= maximo


# --- conversiones ------------------------------------------------------------


def a_json(valor):
    """El valor leído de una celda, listo para guardarse como JSON (las fechas, en ISO)."""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if valor is None or isinstance(valor, (str, int, float, bool)):
        return valor
    return str(valor)


def entero_sector(valor):
    try:
        return int(Decimal(str(valor).strip()))
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"Sector no válido en la planilla: «{valor}».") from error


def texto_de(valor):
    """Texto sin espacios en los bordes; None si está vacío."""
    if valor is None:
        return None
    return str(valor).strip() or None


def numero_de(valor, campo):
    """Devuelve (Decimal o None, texto no válido o None). Vacío es None sin aviso.

    No es válido lo que no es número ni cabe en el campo decimal del modelo.
    """
    texto = texto_de(valor)
    if texto is None:
        return None, None
    try:
        numero = Decimal(texto.replace(",", "."))
    except InvalidOperation:
        return None, texto
    if isinstance(valor, bool) or not numero.is_finite() or not cabe(numero, campo):
        return None, texto
    return numero, None


def cabe(numero, campo):
    """True si el número cabe en el DecimalField del modelo (los campos sin modelo, siempre)."""
    campo_modelo = CAMPOS_DECIMALES.get(campo)
    if campo_modelo is None:
        return True
    try:
        redondeado = numero.quantize(Decimal(1).scaleb(-campo_modelo.decimal_places))
    except InvalidOperation:
        return False
    return len(redondeado.as_tuple().digits) <= campo_modelo.max_digits


def fecha_de(valor):
    """Devuelve (date o None, texto no válido o None). Vacío es None sin aviso."""
    if isinstance(valor, datetime):
        return valor.date(), None
    if isinstance(valor, date):
        return valor, None
    texto = texto_de(valor)
    if texto is None:
        return None, None
    for formato in FORMATOS_FECHA_PLANILLA:
        try:
            return datetime.strptime(texto, formato).date(), None
        except ValueError:
            continue
    return None, texto
