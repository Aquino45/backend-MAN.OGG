"""Carga de la planilla de la brigada en tres pasos que comparten la consola y el panel.

`leer` abre el .xlsx en memoria y devuelve las filas tal como vienen. `evaluar` les aplica la regla
de calidad (`calidad_registro`) y busca duplicados, sin escribir nada. `aplicar_decisiones` guarda
las filas que entran. Entre un paso y otro, el panel guarda las filas del lote (`FilaCarga`) para
preguntar fila por fila; el archivo nunca se guarda.
"""

import hashlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from django.contrib.gis.geos import Point
from django.db import models, transaction
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from apps.censo.models import Arbol, FilaCarga, LoteCarga, Sector, ValorCatalogo
from apps.censo.servicios.calidad_registro import (
    Motivo,
    aplicar,
    aviso_catalogo_sin_cargar,
    buscar_duplicado_por_ubicacion,
    cargar_catalogo,
)
from config.constantes import (
    CAMPOS_CATALOGADOS,
    CAMPOS_CON_NOTA_PLANILLA,
    CAMPOS_DE_UBICACION_PLANILLA,
    CAMPOS_NUMERICOS_PLANILLA,
    CLASIFICACION_CAMPOS_VACIOS,
    CLASIFICACION_DUPLICADO,
    CLASIFICACION_INCONSISTENTE,
    CLASIFICACION_LISTO,
    CLASIFICACIONES_CON_PROBLEMA,
    CO2_PLANILLA_ES_SIMULADO,
    COLUMNAS_PLANILLA,
    DECISION_CONSERVAR,
    DECISION_DESCARTAR,
    ESTADO_LOTE_CONFIRMADO,
    ESTADO_LOTE_VALIDADO,
    ESTADO_REGISTRO_EN_REVISION,
    ETIQUETAS_CAMPOS_CALIDAD,
    FILA_ENCABEZADOS_PLANILLA,
    FORMATO_CODIGO_CALCULADO,
    FORMATO_CODIGO_EXISTENTE_NO_ACTUALIZADO,
    FORMATO_CODIGO_NUEVO_NO_INCORPORADO,
    FORMATO_NOTA_PLANILLA,
    FORMATOS_FECHA_PLANILLA,
    HOJA_PLANILLA,
    MENSAJE_CARGA_LOTE_CERRADO,
    MENSAJE_CARGA_MUCHAS_FILAS,
    MENSAJE_CODIGO_EN_PAPELERA,
    MOTIVO_DUPLICADO_CODIGO,
    MOTIVO_DUPLICADO_UBICACION,
    MOTIVO_FALTA_DATO,
    PATRON_CODIGO_ARBOL,
    SEPARADOR_DE_LINEAS,
    SEPARADOR_NOTAS,
    SRID_UTM_CENSO,
    TIPO_CAMPO_LISTA,
    TIPO_CAMPO_NUMERO,
    TIPO_CAMPO_TEXTO,
)

CAMPOS_SIN_COPIA_DIRECTA = frozenset(
    {"codigo", "sector", "fecha_registro", "observaciones", "control_registro"}
)
"""Campos con tratamiento propio; `control_registro` se lee pero no se guarda."""

CAMPOS_TEXTO = (
    frozenset(COLUMNAS_PLANILLA.values()) - CAMPOS_NUMERICOS_PLANILLA - CAMPOS_SIN_COPIA_DIRECTA
)
"""Campos de texto que se copian tal cual (sin espacios en los bordes) al modelo."""

CAMPOS_CORREGIBLES = frozenset(COLUMNAS_PLANILLA.values()) - {
    "codigo",
    "sector",
    "control_registro",
}
"""Campos que la pantalla de revisión deja corregir."""

CAMPOS_DECIMALES = {
    campo.name: campo
    for campo in Arbol._meta.get_fields()
    if isinstance(campo, models.DecimalField)
}
"""Campos decimales del modelo: de ellos salen los límites de dígitos (`max_digits`, etc.)."""

LARGO_NOMBRE_ARCHIVO = LoteCarga._meta.get_field("nombre_archivo").max_length
"""Cuántos caracteres del nombre del archivo caben en la bitácora."""

ENCABEZADO_POR_CAMPO = {campo: encabezado for encabezado, campo in COLUMNAS_PLANILLA.items()}
"""Inverso de COLUMNAS_PLANILLA, para nombrar la columna en los reportes."""


@dataclass
class ValorNoValido:
    """Un valor de la planilla que no se pudo leer como se esperaba."""

    codigo: str
    columna: str
    texto: str


@dataclass
class ArbolFueraDeSector:
    """Un árbol cuyo punto cae fuera del polígono de su sector."""

    codigo: str
    sector: int
    distancia_m: float


@dataclass
class RegistroEnRevision:
    """Un registro que entró a la BD pero no sale en la web, con sus motivos."""

    codigo: str
    motivos: list[str]


@dataclass
class PlanillaLeida:
    """Las filas de la hoja como vienen: `[(número de fila, {campo: valor})]`."""

    nombre_archivo: str
    filas: list
    catalogos: dict = field(default_factory=dict)


@dataclass
class FilaEvaluada:
    """Una fila ya evaluada: lo que se guardaría, sus motivos y en qué grupo cae."""

    numero_fila: int
    codigo: str
    datos: dict
    valores: dict
    textos_no_numericos: dict
    motivos: list
    clasificacion: str
    codigo_calculado: bool = False
    repetida: bool = False
    duplicado: Motivo | None = None
    duplicado_por_ubicacion: bool = False
    sin_coordenadas: bool = False
    numeros_no_validos: list = field(default_factory=list)
    fechas_no_validas: list = field(default_factory=list)
    fuera_de_catalogo: list = field(default_factory=list)
    fuera_de_sector: list = field(default_factory=list)

    @property
    def motivos_a_mostrar(self):
        return [*self.motivos, *([self.duplicado] if self.duplicado else [])]


@dataclass
class Evaluacion:
    """Todas las filas evaluadas; no se escribió nada."""

    nombre_archivo: str
    filas: list = field(default_factory=list)
    catalogo_calidad: dict = field(default_factory=dict)
    catalogo_sin_cargar: str | None = None

    def cuentas(self):
        """Cuántas filas hay de cada clasificación."""
        return {c: sum(f.clasificacion == c for f in self.filas) for c in CLASIFICACIONES_TODAS}

    def con_problema(self):
        return [f for f in self.filas if f.clasificacion in CLASIFICACIONES_CON_PROBLEMA]


CLASIFICACIONES_TODAS = (
    CLASIFICACION_LISTO,
    CLASIFICACION_CAMPOS_VACIOS,
    CLASIFICACION_INCONSISTENTE,
    CLASIFICACION_DUPLICADO,
)


@dataclass
class ResultadoCarga:
    """Qué pasó con cada fila al guardar."""

    creados: int = 0
    actualizados: int = 0
    publicados: list[str] = field(default_factory=list)
    en_revision: list[RegistroEnRevision] = field(default_factory=list)
    no_incorporados: list[str] = field(default_factory=list)
    ejemplo_sin_cambios: list[str] = field(default_factory=list)
    ejemplo_reemplazados: list[str] = field(default_factory=list)

    def como_dict(self):
        """Los códigos de cada grupo, para guardarlos en la bitácora del lote."""
        return {
            "publicados": self.publicados,
            "en_revision": [registro.codigo for registro in self.en_revision],
            "no_incorporados": self.no_incorporados,
            "ejemplo_sin_cambios": self.ejemplo_sin_cambios,
        }


# --- paso 1: leer ------------------------------------------------------------


def leer(origen, nombre_archivo=None, filas_maximas=None):
    """Lee la hoja de la brigada desde una ruta o desde bytes (`bytes` o `BytesIO`).

    Solo mira los valores guardados: no se evalúa ninguna fórmula ni macro. Lanza ValueError con un
    mensaje claro si el archivo no sirve o declara más de `filas_maximas` filas.
    """
    if isinstance(origen, (str, Path)):
        ruta = Path(origen)
        if not ruta.is_file():
            raise ValueError(f"No existe el archivo: {ruta}")
        contenido, nombre = ruta.read_bytes(), nombre_archivo or ruta.name
    else:
        contenido = origen.getvalue() if isinstance(origen, BytesIO) else bytes(origen)
        nombre = nombre_archivo or ""
    hoja_formulas, hoja_valores = _abrir_hojas(contenido, nombre)
    declaradas = hoja_formulas.max_row - FILA_ENCABEZADOS_PLANILLA
    if filas_maximas is not None and declaradas > filas_maximas:
        raise ValueError(MENSAJE_CARGA_MUCHAS_FILAS.format(filas=declaradas, maximo=filas_maximas))
    columnas = _ubicar_columnas(hoja_formulas)
    return PlanillaLeida(
        nombre_archivo=nombre,
        filas=_leer_filas(hoja_formulas, hoja_valores, columnas),
        catalogos=_leer_catalogos(hoja_formulas, columnas),
    )


def _abrir_hojas(contenido, nombre):
    """Abre el libro dos veces: con fórmulas (y validaciones) y con los valores guardados."""
    try:
        libro_formulas = load_workbook(BytesIO(contenido), data_only=False)
        libro_valores = load_workbook(BytesIO(contenido), data_only=True)
    except Exception as error:  # openpyxl lanza tipos distintos según qué esté roto
        raise ValueError(f"No se pudo abrir {nombre} como planilla de Excel: {error}") from error
    if HOJA_PLANILLA not in libro_formulas.sheetnames:
        raise ValueError(f"La planilla no tiene la hoja «{HOJA_PLANILLA}».")
    return libro_formulas[HOJA_PLANILLA], libro_valores[HOJA_PLANILLA]


def _ubicar_columnas(hoja):
    """Devuelve {campo: número de columna} buscando los encabezados por nombre."""
    posicion_por_encabezado = {
        str(celda.value).strip(): celda.column
        for celda in hoja[FILA_ENCABEZADOS_PLANILLA]
        if celda.value is not None
    }
    faltan = [e for e in COLUMNAS_PLANILLA if e not in posicion_por_encabezado]
    if faltan:
        raise ValueError("Faltan columnas en la planilla: " + ", ".join(faltan))
    return {campo: posicion_por_encabezado[enc] for enc, campo in COLUMNAS_PLANILLA.items()}


def _leer_catalogos(hoja, columnas):
    """Lee las listas de las validaciones de datos: {campo: {valores permitidos}}."""
    campo_por_columna = {columna: campo for campo, columna in columnas.items()}
    catalogos = {}
    for validacion in hoja.data_validations.dataValidation:
        formula = (validacion.formula1 or "").strip()
        if validacion.type != "list" or not (formula.startswith('"') and formula.endswith('"')):
            continue  # solo se leen las listas escritas dentro de la validación
        valores = {valor.strip() for valor in formula.strip('"').split(",") if valor.strip()}
        for rango in str(validacion.sqref).split():
            primera_columna, _, ultima_columna, _ = range_boundaries(rango)
            for columna in range(primera_columna, ultima_columna + 1):
                if columna in campo_por_columna:
                    catalogos.setdefault(campo_por_columna[columna], set()).update(valores)
    return catalogos


def _leer_filas(hoja_formulas, hoja_valores, columnas):
    """Devuelve [(número de fila, {campo: valor guardado})] de las filas que tienen Sector."""
    filas = []
    for numero_fila in range(FILA_ENCABEZADOS_PLANILLA + 1, hoja_formulas.max_row + 1):
        fila = {
            campo: hoja_valores.cell(row=numero_fila, column=columna).value
            for campo, columna in columnas.items()
        }
        if _texto(fila["sector"]) is not None:
            filas.append((numero_fila, fila))
    return filas


# --- paso 2: evaluar ------------------------------------------------------------


def evaluar(planilla):
    """Evalúa todas las filas con la regla de calidad y la búsqueda de duplicados; no escribe.

    Una fila es «listo», «con campos vacíos» (solo le faltan datos), «inconsistente» (algún dato
    no sirve) o «duplicado» (su código se repite en el archivo, o comparte ubicación y especie con
    otro árbol). Lanza ValueError si falta un sector, un código no cumple el patrón o está en la
    papelera.
    """
    catalogo = cargar_catalogo()
    evaluacion = Evaluacion(
        nombre_archivo=planilla.nombre_archivo,
        catalogo_calidad=catalogo,
        catalogo_sin_cargar=aviso_catalogo_sin_cargar(catalogo),
    )
    sectores = _cargar_sectores(planilla.filas)
    orden_en_sector, por_codigo, por_ubicacion = {}, {}, {}
    for numero_fila, datos in planilla.filas:
        id_sector = _entero_sector(datos["sector"])
        orden_en_sector[id_sector] = orden_en_sector.get(id_sector, 0) + 1
        fila = _evaluar_fila(
            numero_fila, datos, planilla, sectores[id_sector], orden_en_sector, catalogo
        )
        _marcar_duplicados(fila, por_codigo, por_ubicacion)
        evaluacion.filas.append(fila)
    return evaluacion


def _cargar_sectores(filas):
    """Trae de la BD los sectores que usa la planilla; falla si falta alguno."""
    ids = {_entero_sector(fila["sector"]) for _, fila in filas}
    sectores = Sector.objects.in_bulk(ids)
    faltan = sorted(ids - set(sectores))
    if faltan:
        raise ValueError(
            "La planilla usa sectores que no están en la BD: "
            + ", ".join(str(id_sector) for id_sector in faltan)
            + ". Importa antes los sectores (importar_sectores)."
        )
    return sectores


def _evaluar_fila(numero_fila, datos, planilla, sector, orden_en_sector, catalogo):
    """Convierte una fila en lo que se guardaría y le aplica la regla de calidad."""
    codigo = _texto(datos["codigo"])
    calculado = codigo is None
    if calculado:
        codigo = FORMATO_CODIGO_CALCULADO.format(sector=sector.id, orden=orden_en_sector[sector.id])
    if not re.fullmatch(PATRON_CODIGO_ARBOL, codigo):
        raise ValueError(f"Fila {numero_fila}: el código «{codigo}» no cumple el patrón.")
    if Arbol.todos.filter(codigo=codigo, eliminado_en__isnull=False).exists():
        raise ValueError(f"Fila {numero_fila}: " + MENSAJE_CODIGO_EN_PAPELERA.format(codigo=codigo))
    fila = FilaEvaluada(
        numero_fila, codigo, datos, {}, {}, [], CLASIFICACION_LISTO, codigo_calculado=calculado
    )
    valores = fila.valores
    valores.update({campo: _texto(datos[campo]) for campo in CAMPOS_TEXTO})
    for campo in valores.keys() & planilla.catalogos.keys():
        if valores[campo] is not None and valores[campo] not in planilla.catalogos[campo]:
            fila.fuera_de_catalogo.append(
                ValorNoValido(codigo, ENCABEZADO_POR_CAMPO[campo], valores[campo])
            )
    notas = _leer_numeros(fila)
    valores["fecha_registro"], texto_fecha = _fecha(datos["fecha_registro"])
    if texto_fecha is not None:
        encabezado = ENCABEZADO_POR_CAMPO["fecha_registro"]
        fila.fechas_no_validas.append(ValorNoValido(codigo, encabezado, texto_fecha))
    partes_observaciones = [_texto(datos["observaciones"]), *notas]
    valores["observaciones"] = SEPARADOR_NOTAS.join(p for p in partes_observaciones if p) or None
    _poner_ubicacion(fila, sector)
    valores["sector"] = sector
    partes_fuente = [
        valores["brigada"],
        f"hoja {HOJA_PLANILLA} de {planilla.nombre_archivo}",
        f"fila {numero_fila}",
    ]
    valores["fuente"] = ", ".join(parte for parte in partes_fuente if parte)
    valores["fila_origen"] = numero_fila
    valores["demo"] = False
    valores["co2_simulado"] = CO2_PLANILLA_ES_SIMULADO
    candidato = Arbol(codigo=codigo, **valores)
    motivos = aplicar(candidato, catalogo, sector.geom, fila.textos_no_numericos)
    fila.motivos, fila.clasificacion = motivos, _clasificar(motivos)
    return fila


def _leer_numeros(fila):
    """Pone en `fila.valores` los números; devuelve las notas para `observaciones`."""
    notas = []
    for campo in (c for c in COLUMNAS_PLANILLA.values() if c in CAMPOS_NUMERICOS_PLANILLA):
        fila.valores[campo], texto = _numero(fila.datos[campo], campo)
        if texto is not None:
            encabezado = ENCABEZADO_POR_CAMPO[campo]
            fila.numeros_no_validos.append(ValorNoValido(fila.codigo, encabezado, texto))
            fila.textos_no_numericos[campo] = texto
            if campo in CAMPOS_CON_NOTA_PLANILLA:
                notas.append(FORMATO_NOTA_PLANILLA.format(encabezado=encabezado, texto=texto))
    return notas


def _poner_ubicacion(fila, sector):
    valores = fila.valores
    este, norte = valores.pop("utm_este_m"), valores.pop("utm_norte_m")
    valores["ubicacion"] = None
    if este is not None and norte is not None:
        valores["ubicacion"] = Point(float(este), float(norte), srid=SRID_UTM_CENSO)
    else:
        fila.sin_coordenadas = True
    if valores["ubicacion"] is not None and not sector.geom.covers(valores["ubicacion"]):
        distancia_m = sector.geom.distance(valores["ubicacion"])
        fila.fuera_de_sector.append(ArbolFueraDeSector(fila.codigo, sector.id, distancia_m))


def _clasificar(motivos):
    """Sin motivos es listo; si solo faltan datos, campos vacíos; si no, inconsistente."""
    if not motivos:
        return CLASIFICACION_LISTO
    if all(_es_falta_de_dato(motivo) for motivo in motivos):
        return CLASIFICACION_CAMPOS_VACIOS
    return CLASIFICACION_INCONSISTENTE


def _es_falta_de_dato(motivo):
    etiqueta = ETIQUETAS_CAMPOS_CALIDAD.get(motivo.campo)
    return etiqueta is not None and motivo.porque == MOTIVO_FALTA_DATO.format(etiqueta=etiqueta)


def _marcar_duplicados(fila, por_codigo, por_ubicacion):
    """Marca como duplicadas las filas con código repetido o con la ubicación y especie de otra."""
    if fila.codigo in por_codigo:
        primera = por_codigo[fila.codigo]
        fila.repetida = True
        for afectada in (primera, fila):
            afectada.clasificacion = CLASIFICACION_DUPLICADO
            afectada.duplicado = afectada.duplicado or Motivo(
                "codigo", MOTIVO_DUPLICADO_CODIGO.format(codigo=fila.codigo)
            )
        return
    por_codigo[fila.codigo] = fila
    ubicacion, especie = fila.valores["ubicacion"], fila.valores["nombre_cientifico"]
    clave = (ubicacion.x, ubicacion.y, especie.lower()) if ubicacion and especie else None
    otro = por_ubicacion.get(clave) or buscar_duplicado_por_ubicacion(
        fila.codigo, ubicacion, especie
    )
    if otro:
        fila.duplicado_por_ubicacion = True
        fila.clasificacion = CLASIFICACION_DUPLICADO
        fila.duplicado = Motivo("ubicacion", MOTIVO_DUPLICADO_UBICACION.format(otro=otro))
    elif clave:
        por_ubicacion[clave] = fila.codigo


# --- paso 3: aplicar las decisiones ------------------------------------------------------------


def aplicar_decisiones(evaluacion, decisiones):
    """Guarda las filas que entran y devuelve qué pasó con cada una.

    `decisiones` es `{número de fila: conservar | descartar}`. Una fila listo o con campos vacíos
    entra sin decisión; una inconsistente o duplicada solo si se conserva. Todo lo que se guarda
    pasa por la regla de calidad, así que entra publicado o en revisión.
    """
    resultado = ResultadoCarga()
    with transaction.atomic():
        for fila in evaluacion.filas:
            decision = decisiones.get(fila.numero_fila)
            sin_problema = fila.clasificacion not in CLASIFICACIONES_CON_PROBLEMA
            if decision == DECISION_CONSERVAR or (decision is None and sin_problema):
                _guardar_fila(fila, evaluacion.catalogo_calidad, resultado)
            else:
                resultado.no_incorporados.append(fila.codigo)
    return resultado


def _guardar_fila(fila, catalogo, resultado):
    """Guarda la fila con su estado. Un árbol de ejemplo solo se reemplaza por una fila completa.

    Un duplicado de ubicación que se conserva entra siempre en revisión, con ese motivo.
    """
    codigo = fila.codigo
    candidato = Arbol(codigo=codigo, **fila.valores)
    motivos = aplicar(candidato, catalogo, fila.valores["sector"].geom, fila.textos_no_numericos)
    if fila.duplicado_por_ubicacion:
        motivos = [*motivos, fila.duplicado]
        candidato.estado_registro = ESTADO_REGISTRO_EN_REVISION
        candidato.motivos_revision = SEPARADOR_DE_LINEAS.join(str(m) for m in motivos)
    existente = Arbol.todos.filter(codigo=codigo).first()
    if existente is not None and existente.demo:
        if motivos:
            resultado.ejemplo_sin_cambios.append(codigo)
            return
        resultado.ejemplo_reemplazados.append(codigo)
    estado = {
        "estado_registro": candidato.estado_registro,
        "motivos_revision": candidato.motivos_revision,
        "co2_simulado": candidato.co2_simulado,
    }
    _, creado = Arbol.todos.update_or_create(codigo=codigo, defaults={**fila.valores, **estado})
    if creado:
        resultado.creados += 1
    else:
        resultado.actualizados += 1
    if motivos:
        resultado.en_revision.append(RegistroEnRevision(codigo, [str(m) for m in motivos]))
    else:
        resultado.publicados.append(codigo)


# --- el lote del panel ------------------------------------------------------------


def valores_activos_del_catalogo():
    """`{campo: [valores activos]}` en el orden del catálogo, para las listas de la revisión."""
    activos = {campo: [] for campo in CAMPOS_CATALOGADOS}
    filas = ValorCatalogo.objects.filter(activo=True).order_by("campo", "orden", "valor")
    for campo, valor in filas.values_list("campo", "valor"):
        activos.setdefault(campo, []).append(valor)
    return activos


def registrar_lote(usuario, nombre_archivo, contenido, evaluacion):
    """Guarda el lote (con el hash del archivo, no el archivo) y una fila por registro leído."""
    cuentas = evaluacion.cuentas()
    with transaction.atomic():
        lote = LoteCarga.objects.create(
            usuario=usuario,
            nombre_archivo=nombre_archivo[:LARGO_NOMBRE_ARCHIVO],
            sha256=hashlib.sha256(contenido).hexdigest(),
            filas_leidas=len(evaluacion.filas),
            listos=cuentas[CLASIFICACION_LISTO],
            con_campos_vacios=cuentas[CLASIFICACION_CAMPOS_VACIOS],
            inconsistentes=cuentas[CLASIFICACION_INCONSISTENTE],
            duplicados=cuentas[CLASIFICACION_DUPLICADO],
        )
        FilaCarga.objects.bulk_create(
            FilaCarga(
                lote=lote,
                numero_fila=fila.numero_fila,
                codigo=fila.codigo,
                clasificacion=fila.clasificacion,
                datos={campo: _a_json(valor) for campo, valor in fila.datos.items()},
                motivos=[{"campo": m.campo, "porque": m.porque} for m in fila.motivos_a_mostrar],
            )
            for fila in evaluacion.filas
        )
    return lote


def filas_con_problema(lote):
    """Las filas del lote por las que hay que preguntar, en el orden de la planilla."""
    return lote.filas.filter(clasificacion__in=CLASIFICACIONES_CON_PROBLEMA)


def mensajes_de_no_incorporacion(lote):
    """Qué pasa, por código, si se confirma sin revisar: un mensaje distinto si ya existe."""
    codigos = list(dict.fromkeys(filas_con_problema(lote).values_list("codigo", flat=True)))
    existentes = set(Arbol.todos.filter(codigo__in=codigos).values_list("codigo", flat=True))
    return [
        (
            FORMATO_CODIGO_EXISTENTE_NO_ACTUALIZADO
            if codigo in existentes
            else FORMATO_CODIGO_NUEVO_NO_INCORPORADO
        ).format(codigo=codigo)
        for codigo in codigos
    ]


def campos_a_corregir(fila, activos):
    """Los campos de una fila con problema que se pueden editar: `[{campo, etiqueta, ...}]`."""
    campos = []
    for motivo in fila.motivos:
        pedidos = (
            CAMPOS_DE_UBICACION_PLANILLA if motivo["campo"] == "ubicacion" else (motivo["campo"],)
        )
        campos += [c for c in pedidos if c in CAMPOS_CORREGIBLES and c not in campos]
    return [_describir_campo(campo, fila.datos.get(campo), activos) for campo in campos]


def _describir_campo(campo, valor, activos):
    if campo in CAMPOS_NUMERICOS_PLANILLA:
        tipo = TIPO_CAMPO_NUMERO
    elif activos.get(campo):
        tipo = TIPO_CAMPO_LISTA
    else:
        tipo = TIPO_CAMPO_TEXTO
    return {
        "campo": campo,
        "etiqueta": ENCABEZADO_POR_CAMPO[campo],
        "tipo": tipo,
        "valor_en_planilla": _texto(valor),
        "valor": _valor_mostrado(campo, valor, tipo),
        "opciones": activos.get(campo, []) if tipo == TIPO_CAMPO_LISTA else [],
    }


def _valor_mostrado(campo, valor, tipo):
    """Lo que trae el control de la pantalla; un texto donde va un número se muestra vacío."""
    if tipo == TIPO_CAMPO_NUMERO and _numero(valor, campo)[0] is None:
        return ""
    return _texto(valor) or ""


def ediciones_del_formulario(filas, activos, pedidos):
    """`{pk de la fila: {campo: valor}}` con lo que el usuario cambió; el resto se ignora.

    `pedidos` es una función `(pk, campo) -> texto o None` que lee el formulario. Solo cuentan los
    campos que la pantalla ofrece y que cambiaron frente a lo que se mostraba.
    """
    ediciones = {}
    for fila in filas:
        for descrito in campos_a_corregir(fila, activos):
            nuevo = pedidos(fila.pk, descrito["campo"])
            if nuevo is not None and nuevo.strip() != descrito["valor"]:
                ediciones.setdefault(fila.pk, {})[descrito["campo"]] = nuevo.strip()
    return ediciones


def codigo_con_varias_conservadas(filas, conservadas):
    """El primer código con más de una fila conservada, o None (se conserva como mucho una)."""
    vistos = set()
    for fila in filas:
        if fila.pk in conservadas:
            if fila.codigo in vistos:
                return fila.codigo
            vistos.add(fila.codigo)
    return None


def confirmar_lote(lote, conservadas=None, ediciones=None):
    """Evalúa otra vez el lote (con lo corregido), guarda lo que entra y lo cierra.

    Con `conservadas` (los pk de las filas elegidas) las demás filas con problema no entran;
    sin él, no entra ninguna. Las filas buenas entran siempre. Las filas del lote se borran.
    Lanza ValueError si el lote ya no está pendiente.
    """
    ediciones = ediciones or {}
    with transaction.atomic():
        lote = LoteCarga.objects.select_for_update().get(pk=lote.pk)
        if lote.estado != ESTADO_LOTE_VALIDADO:
            raise ValueError(MENSAJE_CARGA_LOTE_CERRADO)
        filas = list(lote.filas.all())
        planilla = PlanillaLeida(
            lote.nombre_archivo,
            [(f.numero_fila, {**f.datos, **ediciones.get(f.pk, {})}) for f in filas],
        )
        evaluacion = evaluar(planilla)
        if conservadas is None:
            decisiones = {f.numero_fila: DECISION_DESCARTAR for f in evaluacion.con_problema()}
        else:
            decisiones = {
                f.numero_fila: DECISION_CONSERVAR if f.pk in conservadas else DECISION_DESCARTAR
                for f in filas
                if f.clasificacion in CLASIFICACIONES_CON_PROBLEMA
            }
        for fila in filas:
            fila.decision = decisiones.get(fila.numero_fila, fila.decision)
        FilaCarga.objects.bulk_update(filas, ["decision"])
        resultado = aplicar_decisiones(evaluacion, decisiones)
        lote.cerrar(ESTADO_LOTE_CONFIRMADO, resultado.como_dict())
    return resultado


# --- conversiones ------------------------------------------------------------


def _a_json(valor):
    """El valor leído de una celda, listo para guardarse como JSON (las fechas, en ISO)."""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if valor is None or isinstance(valor, (str, int, float, bool)):
        return valor
    return str(valor)


def _entero_sector(valor):
    try:
        return int(Decimal(str(valor).strip()))
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"Sector no válido en la planilla: «{valor}».") from error


def _texto(valor):
    """Texto sin espacios en los bordes; None si está vacío."""
    if valor is None:
        return None
    return str(valor).strip() or None


def _numero(valor, campo):
    """Devuelve (Decimal o None, texto no válido o None). Vacío es None sin aviso.

    No es válido lo que no es número ni cabe en el campo decimal del modelo.
    """
    texto = _texto(valor)
    if texto is None:
        return None, None
    try:
        numero = Decimal(texto.replace(",", "."))
    except InvalidOperation:
        return None, texto
    if isinstance(valor, bool) or not numero.is_finite() or not _cabe(numero, campo):
        return None, texto
    return numero, None


def _cabe(numero, campo):
    """True si el número cabe en el DecimalField del modelo (los campos sin modelo, siempre)."""
    campo_modelo = CAMPOS_DECIMALES.get(campo)
    if campo_modelo is None:
        return True
    try:
        redondeado = numero.quantize(Decimal(1).scaleb(-campo_modelo.decimal_places))
    except InvalidOperation:
        return False
    return len(redondeado.as_tuple().digits) <= campo_modelo.max_digits


def _fecha(valor):
    """Devuelve (date o None, texto no válido o None). Vacío es None sin aviso."""
    if isinstance(valor, datetime):
        return valor.date(), None
    if isinstance(valor, date):
        return valor, None
    texto = _texto(valor)
    if texto is None:
        return None, None
    for formato in FORMATOS_FECHA_PLANILLA:
        try:
            return datetime.strptime(texto, formato).date(), None
        except ValueError:
            continue
    return None, texto
